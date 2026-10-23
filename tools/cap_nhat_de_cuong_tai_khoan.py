"""Update the Java LAN proposal in place, retaining a byte-identical backup."""
from copy import deepcopy
from datetime import datetime, timezone
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import json
import math

from docx import Document
from docx.text.paragraph import Paragraph
from PIL import Image, ImageDraw, ImageFont

from cap_nhat_de_cuong_java_lan import replace_paragraph, replace_cell, replace_table

ROOT = Path(__file__).resolve().parents[1]
TARGET = ROOT / "output/documents/De_cuong_Network_Quiz_Battle_Java_LAN_Nop_Giang_Vien.docx"
TMP = ROOT / "tmp/de_cuong_tai_khoan"
REFERENCE = TMP / "reference.docx"
EXPECTED_HASH = "1c7e1ec4f96564b599d8e3e578b95a0aeb46d18c6199c750dd11f50ca086e54e"


def insert_after(anchor, template, text, label=False):
    element = deepcopy(template._p)
    anchor._p.addnext(element)
    result = Paragraph(element, anchor._parent)
    replace_paragraph(result, text, label)
    return result


def diagram_bytes():
    img = Image.new("RGB", (1700, 700), "white")
    draw = ImageDraw.Draw(img)
    fonts = Path(r"C:\Windows\Fonts")
    bold = ImageFont.truetype(str(fonts / "arialbd.ttf"), 29)
    body = ImageFont.truetype(str(fonts / "arial.ttf"), 25)
    small = ImageFont.truetype(str(fonts / "arial.ttf"), 23)

    def box(xy, title, lines):
        x1, y1, x2, y2 = xy
        draw.rounded_rectangle(xy, radius=20, fill="#EAF3F9", outline="#1F6FAB", width=4)
        for i, (line, font) in enumerate([(title, bold)] + [(line, body) for line in lines]):
            width = draw.textlength(line, font=font)
            assert width < x2 - x1 - 24, line
            draw.text(((x1 + x2 - width) / 2, y1 + 23 + i * 37), line, font=font, fill="#17365D")

    def arrow(a, b):
        draw.line([a, b], fill="#1F6FAB", width=7)
        for start, end in [(a, b), (b, a)]:
            angle = math.atan2(end[1] - start[1], end[0] - start[0])
            points = [end] + [(end[0] - 19 * math.cos(angle + d), end[1] - 19 * math.sin(angle + d)) for d in [-0.52, 0.52]]
            draw.polygon(points, fill="#1F6FAB")

    box((55, 35, 450, 180), "GIẢNG VIÊN VÀ ADMIN", ["JavaFX và đăng nhập", "Ngân hàng riêng, phòng thi"])
    box((55, 220, 450, 365), "HỌC SINH 1", ["JavaFX và mã phòng", "Nickname, ready, trả lời"])
    box((55, 405, 450, 550), "HỌC SINH 2 ĐẾN 10", ["Cùng phòng thi", "Nhận điểm từ Server"])
    box((610, 160, 975, 420), "MẠNG LAN", ["Switch hoặc Wi-Fi", "TCP Socket qua TLS", "JSON UTF-8 từng dòng", "Cổng mặc định 5050"])
    box((1130, 55, 1660, 335), "JAVA SERVER", ["SSLServerSocket và worker", "Xác thực, kiểm tra quyền", "Game Worker và deadline", "Chấm điểm và xếp hạng"])
    box((1165, 435, 1625, 580), "SQLITE QUA JDBC", ["Tài khoản, câu hỏi, kết quả", "Owner ID theo tài khoản"])
    for y in [107, 292, 477]:
        arrow((450, y), (610, 290))
    arrow((975, 290), (1130, 197))
    arrow((1395, 335), (1395, 435))
    draw.text((65, 605), "Ví dụ: Server 192.168.1.10:5050; các Client cùng LAN, xác minh chứng chỉ trước khi gửi dữ liệu.", font=small, fill="#17365D")
    draw.text((65, 645), "Chỉ Server đọc ghi SQLite; dữ liệu gắn với tài khoản trên Server đó. Không cần Internet hoặc hosting.", font=small, fill="#17365D")
    buf = BytesIO()
    img.save(buf, format="PNG", dpi=(220, 220))
    return buf.getvalue()


def main():
    TMP.mkdir(parents=True, exist_ok=True)
    if not REFERENCE.exists():
        source_bytes = TARGET.read_bytes()
        assert sha256(source_bytes).hexdigest() == EXPECTED_HASH, "Source changed; inspect user edits before proceeding"
        REFERENCE.write_bytes(source_bytes)
    source_bytes = REFERENCE.read_bytes()
    assert sha256(source_bytes).hexdigest() == EXPECTED_HASH
    with ZipFile(BytesIO(source_bytes)) as z:
        original_parts = {p.filename: z.read(p.filename) for p in z.infolist()}
    doc = Document(BytesIO(source_bytes))
    blocks = list(doc.iter_inner_content())
    tables = list(doc.tables)
    assert len(blocks) == 122 and len(tables) == 13
    updates = {
        4: "Kết quả chính: Server Java trong LAN; Client JavaFX dành cho giảng viên, Admin và học sinh kết nối TCP qua TLS. SQLite lưu tài khoản, ngân hàng câu hỏi riêng và kết quả tại máy Server.",
        9: "Vấn đề cần giải quyết: Server phục vụ nhiều Client, đồng bộ câu hỏi và thời hạn, ngăn nộp lặp và trả bảng điểm nhất quán. Giảng viên đăng nhập để sử dụng ngân hàng riêng và tổ chức ván thi qua mạng.",
        11: "Đối tượng sử dụng: Admin hệ thống quản lý tài khoản giảng viên; giảng viên có ngân hàng riêng và điều hành phòng mình tạo. Học sinh tham gia bằng nickname và mã phòng, chưa cần tài khoản. Server kiểm tra quyền cho từng yêu cầu.",
        12: "Sản phẩm dự kiến: Server console, Client quản lý cho giảng viên/Admin, Client học sinh, ít nhất 20 câu hỏi mẫu, hướng dẫn TLS và chạy LAN, bộ kiểm thử, video demo và báo cáo đồ án.",
        15: "Một Server Java trong LAN; Client quản lý và Client học sinh viết bằng JavaFX, sử dụng TCP Socket qua TLS.",
        16: "Đăng nhập, đăng xuất và đổi mật khẩu cho Admin/giảng viên; Admin tạo, khóa và đặt lại mật khẩu tài khoản giảng viên. Học sinh vào bằng nickname/mã phòng; chỉ Server truy cập DB.",
        17: "Thêm, xem, sửa, xóa và tìm câu hỏi bốn lựa chọn trong ngân hàng của tài khoản đang đăng nhập; dữ liệu giữ trong SQLite khi đóng ứng dụng hoặc restart Server.",
        18: "Một phòng tại một thời điểm, có chủ phòng đã đăng nhập; cấu hình tên, số câu, thời lượng và chỉ lấy câu hỏi từ ngân hàng của chủ phòng.",
        21: "Lưu kết quả chung cuộc theo tài khoản tổ chức; chủ phòng xem lịch sử riêng, reset hoặc đóng phòng sau ván thi.",
        23: "Ngoài phạm vi: Hosting/Internet, nhiều phòng đồng thời, tài khoản học sinh, đồng bộ giữa các Server, khôi phục phiên thi, web/mobile, chat/voice, lưu từng đáp án và chống gian lận nâng cao. Nhập CSV và bộ cài là hướng mở rộng.",
        24: "Giả định và giới hạn: Có nhiều tài khoản giảng viên nhưng chỉ một chủ phòng điều khiển một phòng tại một thời điểm. Nghiệm thu tối thiểu 3 học sinh, thử tải đến 10; các máy cùng LAN, tin cậy chứng chỉ Server và truy cập được cổng TCP. Server phải duy trì hoạt động khi thi.",
        34: "Ba học sinh nhận cùng session_id và question_id; chỉ đáp án đầu tiên hợp lệ được nhận, đáp án trễ bị TIME_EXPIRED. Học sinh không gọi được API quản lý; giảng viên không truy cập ngân hàng/phòng của tài khoản khác. Ranking cuối khớp Server và lịch sử đã lưu; lỗi lưu phải được báo.",
        39: "Nguyên tắc: Server giữ trạng thái chuẩn, deadline, đáp án đúng và điểm. Client quản lý phục vụ giảng viên/Admin, luôn gọi API qua Socket dù chạy cùng máy Server. Vai trò và quyền sở hữu lấy từ tài khoản đã xác thực; không tin role hoặc owner_id do Client tự khai.",
        45: "Kênh truyền: TCP Socket qua TLS bằng JSSE, cổng LAN mặc định 5050. Sau xác minh chứng chỉ mới trao đổi JSON UTF-8 một dòng kết thúc newline; receiver ghép byte và chỉ parse frame hoàn chỉnh, tối đa 64 KiB. Không dùng Java object serialization.",
        46: "Cấu trúc chung: protocol_version, type, request_id, room_code và payload. session_id là ID ván, không phải phiên đăng nhập. user_id/role hoặc player_id gắn với connection tại Server; quyền không nằm ở lựa chọn màn hình của Client.",
        57: "1.  Chọn Giảng viên, kết nối TLS và đăng nhập; Server trả user_id/role. Giảng viên CRUD ngân hàng riêng, tạo phòng và trở thành chủ phòng.",
        58: "2.  Chọn Học sinh, gửi JOIN với nickname và mã phòng qua TLS; Server gán player_id, trả danh sách và nhận READY_SET.",
        59: "3.  Chủ phòng gửi START_GAME; Server snapshot câu hỏi của chủ phòng/người chơi, tạo session_id và phát GAME_START.",
        62: "6.  Sau câu cuối, Server lưu kết quả theo host_user_id, phát GAME_OVER; chủ phòng xem lịch sử, reset chơi tiếp hoặc đóng phòng để giảng viên khác tổ chức.",
        64: "Accept thread dùng SSLServerSocket; TLS handshake chạy ở worker có timeout. Mỗi connection có reader và writer duy nhất, queue gửi riêng có giới hạn; Client chậm không chặn broadcast.",
        65: "Một Game Worker sở hữu trạng thái ván và JDBC connection, nhận sự kiện qua BlockingQueue và không ghi socket trực tiếp. PBKDF2 chạy ở auth worker có giới hạn, trả kết quả về Game Worker để không làm trễ khóa câu.",
        67: "PING/PONG mỗi 5 giây, timeout 20 giây. Học sinh rớt mạng giữ điểm/đáp án đã nhận, không khôi phục phiên; người khác tiếp tục. Chủ phòng rớt mạng thì ván vẫn tự chạy, chỉ cùng tài khoản đăng nhập lại mới nhận quyền điều khiển.",
        68: "Hết học sinh kết nối thì Server hủy ván; chủ phòng đăng nhập lại để reset. JSON sai, message lớn, sai quyền, sai owner hoặc sai trạng thái bị từ chối; không làm dừng toàn Server.",
        71: "Lưu trữ: Phòng và chủ phòng, người chơi, ready, deadline và đáp án hiện tại ở RAM. SQLite lưu users, câu hỏi và kết quả theo tài khoản trên Server này. Kết quả lưu bằng transaction; nếu lỗi, GAME_OVER vẫn trả ranking nhưng báo result_saved=false.",
        73: "Mô hình: Một máy chạy Server console; Client giảng viên/Admin và học sinh cùng switch hoặc Wi-Fi. Ví dụ 192.168.1.10:5050; thử một máy dùng 127.0.0.1. Chứng chỉ cần SAN phù hợp với IP/hostname Client kết nối.",
        76: "Bảo vệ dữ liệu: TLS mã hóa kênh truyền trong LAN; PBKDF2 bảo vệ mật khẩu lưu trữ. SQLite không tự được mã hóa: hạn chế quyền truy cập file DB và bản sao lưu. Không dùng mật khẩu cá nhân hoặc mở dịch vụ ra Internet.",
        78: "Màn hình đầu: Hỏi Bạn là ai với hai lựa chọn Giảng viên và Học sinh; lối vào Admin hệ thống riêng. Chọn vai trò chỉ điều hướng, không cấp quyền. Client quản lý có Đăng nhập, Phòng thi, Ngân hàng riêng, Lịch sử và Đổi mật khẩu; tab Tài khoản chỉ hiện cho ADMIN.",
        79: "Client học sinh: IP LAN, cổng, nickname và mã phòng; không có bước đăng ký. Tiếp theo là lobby/ready, bốn nút đáp án, kết quả câu và bảng hạng cuối. Chỉ báo ghi nhận sau ACK; khóa nút khi gửi hoặc hết hạn.",
        81: "Mã nguồn: Giữ 5 module Maven nqb-shared, nqb-server, nqb-client-core, nqb-admin-client và nqb-player-client. nqb-admin-client phục vụ cả giảng viên/Admin, nqb-player-client phục vụ học sinh. TLS và kết nối nằm trong client-core; auth/quyền/repository tại Server; JavaFX chỉ ở hai module UI.",
        90: "Phối hợp: Thiện chủ trì TLS, auth, quyền và room owner; Mỹ làm UI đăng nhập/tài khoản, schema users, ngân hàng/lịch sử theo owner và scoring; Trường Thi làm chọn vai trò/học sinh, client-core và vòng thi. Cả nhóm thống nhất protocol, kiểm thử chéo hai tài khoản và thử nhiều máy LAN.",
        97: "1.  Chạy Server LAN, nạp chứng chỉ/truststore; Admin đăng nhập và tạo hai tài khoản giảng viên demo.",
        98: "2.  Chọn Giảng viên, đăng nhập tài khoản A/B; chứng minh ngân hàng riêng, đổi mật khẩu, lưu qua đăng nhập lại; A tạo phòng từ câu hỏi của mình.",
        99: "3.  Ba học sinh chọn Học sinh, nhập nickname/mã phòng và ready; đối chiếu lobby, thử lệnh quản lý trái quyền.",
        102: "6.  Đối chiếu ranking với lịch sử chủ phòng, reset/đóng phòng; restart Server và đăng nhập lại để kiểm tra ngân hàng/lịch sử vẫn thuộc đúng tài khoản.",
        103: "7.  Wireshark trên traffic được phép: quan sát TCP/TLS handshake và dữ liệu đã mã hóa. Giải thích JSON newline bằng mẫu protocol đã bỏ thông tin nhạy cảm; không kỳ vọng đọc được JSON hoặc mật khẩu trong capture TLS.",
        111: "Không đưa mật khẩu, private key hoặc dữ liệu thật vào repository/log. Cấp tài khoản demo riêng, buộc đổi mật khẩu tạm; hướng dẫn truststore và xác minh chứng chỉ. Bản sao lưu DB chỉ dành cho nhóm được phép; không công bố cổng ra Internet.",
    }
    for i, text in updates.items():
        assert isinstance(blocks[i], Paragraph)
        replace_paragraph(blocks[i], text, label=": " in text and text[:1].isalpha())

    replace_cell(tables[0].cell(4, 1), "Phạm Phúc Thiện - 2380610421 - F01, F02, F04, F05, F12; TLS/TCP Server, tài khoản/quyền, phòng, lỗi và điều phối nhóm.")
    replace_cell(tables[0].cell(5, 1), "Trịnh Viết Mỹ - 2380601382 - F03, F09, F11; Client quản lý, dữ liệu theo tài khoản, điểm và lịch sử.\nNguyễn Trường Thi - 2380602092 - F06, F07, F08, F10; Client học sinh, chọn vai trò và vòng thi.")
    changed_functions = {
        1: ["F01", "Kết nối Server", "Socket/ServerSocket qua TLS; xác minh chứng chỉ trước gửi dữ liệu.", "IP/cổng, truststore / trạng thái", "Sai IP, CA/SAN, timeout"],
        2: ["F02", "Tài khoản và phân quyền", "Login/Logout, đổi mật khẩu; Admin cấp/khóa/reset tài khoản giảng viên.", "Thông tin đăng nhập / user, role", "Sai mật khẩu, bị khóa, sai quyền"],
        3: ["F03", "Ngân hàng câu hỏi riêng", "CRUD/tìm kiếm qua TCP; Server lọc owner_user_id từ tài khoản.", "Câu hỏi / ngân hàng riêng", "Sai dữ liệu, khác chủ sở hữu"],
        4: ["F04", "Cấu hình và chủ phòng", "Một phòng có host_user_id; lấy câu hỏi của chủ phòng khi chờ.", "Cấu hình / ROOM_STATE", "Phòng bận, sai chủ, thiếu câu"],
        6: ["F06", "Sẵn sàng và bắt đầu", "Ready; chỉ chủ phòng Start khi đủ điều kiện, countdown 3 giây.", "Ready, Start / GAME_START", "Thiếu người, chưa ready"],
        11: ["F11", "Kết quả và lịch sử riêng", "Lưu kết quả theo host_user_id; xem lịch sử riêng, reset/đóng phòng.", "Ván / GAME_OVER, lịch sử", "Khác chủ, lỗi DB, reset sớm"],
        12: ["F12", "Giám sát và lỗi", "Heartbeat, TLS/auth timeout, rate limit, queue giới hạn, log an toàn.", "Sự kiện / lỗi rõ ràng", "Mất mạng, JSON sai, quá tải"],
    }
    for row_index, record in changed_functions.items():
        for cell, text in zip(tables[1].rows[row_index].cells, record):
            replace_cell(cell, text)

    replace_table(tables[2], [
        ["Client quản lý", "JavaFX cho giảng viên/Admin; login, ngân hàng riêng, phòng, lịch sử; ADMIN có tab tài khoản."],
        ["Client học sinh", "JavaFX; nickname/mã phòng, lobby, ready, quiz và kết quả; dùng client-core chung."],
        ["TCP/TLS Server", "Java console; SSLServerSocket, handshake/reader/writer worker, không phụ thuộc JavaFX."],
        ["Auth và quyền", "PBKDF2; user_id/role theo connection, kiểm tra owner của câu hỏi, lịch sử và phòng."],
        ["Game Worker", "BlockingQueue điều phối ván, deadline, auth result và JDBC; một chủ sở hữu trạng thái."],
        ["Scoring", "Đúng +100, không thưởng tốc độ, đồng điểm đồng hạng; ranking từ Server."],
        ["Repository SQLite", "users, questions, matches, match_results; PreparedStatement lọc owner và transaction."],
        ["Logging và config", "IP/cổng, DB, keystore/truststore; không log mật khẩu/hash hoặc private key."],
    ])
    replace_table(tables[3], [
        ["Xác thực", "LOGIN, LOGIN_OK, LOGIN_ERROR, LOGOUT, PASSWORD_CHANGE; user_id/role do Server trả"],
        ["Tài khoản", "USER_LIST, USER_CREATE, USER_SET_ACTIVE, USER_RESET_PASSWORD; chỉ ADMIN"],
        ["Ngân hàng", "QUESTION_LIST/CREATE/UPDATE/DELETE, HISTORY_LIST; lọc theo tài khoản đăng nhập"],
        ["Phòng", "ROOM_CONFIGURE, JOIN, JOIN_OK, READY_SET, PLAYER_LIST, ROOM_STATE, LEAVE, ROOM_RESET, ROOM_CLOSE"],
        ["Thi đấu", "START_GAME, GAME_START, QUESTION, ANSWER, ANSWER_ACK, QUESTION_RESULT, ABORT_GAME, GAME_ABORTED"],
        ["Kết quả", "GAME_OVER; ranking gồm player_id, nickname, score, rank, connected; có result_saved"],
        ["Kết nối và lỗi", "PING, PONG, ERROR; phản hồi tương ứng request_id, ván/câu nếu áp dụng"],
    ])
    replace_table(tables[4], [
        ["users", "id, username UNIQUE, display_name, role, password_hash, password_salt, kdf_algorithm, kdf_iterations, credential_version, is_active, must_change_password, created_at_utc"],
        ["questions", "id, owner_user_id FK users, content, option_a..d, correct_answer, category, is_active"],
        ["matches", "session_id, host_user_id FK users, room_name, started_at_utc, ended_at_utc, question_count, duration_seconds, status"],
        ["match_results", "session_id FK matches, player_id, nickname, score, correct_count, rank, connected_at_end; khóa ghép session_id/player_id"],
    ])
    replace_cell(tables[5].cell(3, 1), "Socket/ServerSocket, JSSE")
    replace_cell(tables[5].cell(3, 2), "Tự lập trình TCP stream trong LAN, dùng SSLSocket/SSLServerSocket bảo vệ kênh truyền.")
    replace_cell(tables[5].cell(7, 2), "DB tại Server; khóa ngoại owner, PreparedStatement và transaction kết quả.")
    replace_cell(tables[5].cell(8, 1), "PBKDF2, SecureRandom, TLS")
    replace_cell(tables[5].cell(8, 2), "Băm có salt; role/owner tại Server; chứng chỉ/truststore cho TCP qua TLS [7], [8].")
    replace_cell(tables[6].cell(1, 1), "TCP/IP, Socket, ServerSocket và TLS")
    replace_cell(tables[6].cell(1, 2), "Giải thích accept/connect, TLS handshake, login và quyền theo connection.")
    replace_cell(tables[6].cell(6, 2), "CRUD theo owner bằng PreparedStatement; dữ liệu riêng còn sau restart.")
    replace_cell(tables[6].cell(8, 2), "Phân biệt loopback/IP LAN/bind; kiểm tra cổng, chứng chỉ và capture TCP/TLS.")
    replace_cell(tables[7].cell(2, 1), "Chốt protocol/TLS, schema users/owner và UI; thử Socket hai máy LAN")
    replace_cell(tables[7].cell(3, 1), "Server/Client console, TLS, login/PBKDF2, quyền và chủ phòng")
    replace_cell(tables[7].cell(4, 1), "SQLite, ngân hàng riêng, UI quản lý/tài khoản và dữ liệu mẫu")
    replace_cell(tables[7].cell(6, 1), "Heartbeat, lỗi LAN/TLS, tách dữ liệu hai tài khoản và test nhiều máy")
    replace_cell(tables[8].cell(1, 4), "Java TCP/TLS Server, login/PBKDF2, quyền/owner, phòng, lỗi và tích hợp.")
    replace_cell(tables[8].cell(2, 4), "Client quản lý JavaFX, UI login/tài khoản, SQLite users/owner, CRUD, điểm và lịch sử.")
    replace_cell(tables[8].cell(3, 4), "Client học sinh JavaFX, chọn vai trò, client-core/TLS, vòng câu/đáp án và bảng hạng.")
    risks = [[c.text for c in row.cells] for row in tables[9].rows[1:]]
    risks.insert(2, ["Lộ mật khẩu hoặc truy cập chéo tài khoản", "Cao", "PBKDF2/salt, TLS xác minh CA/SAN, kiểm tra role/owner tại Server; test hai giảng viên và log an toàn."])
    risks.insert(3, ["Sai chứng chỉ hoặc auth làm trễ vòng thi", "Trung bình", "Chuẩn bị IP/SAN và truststore sớm; PBKDF2/handshake ở worker giới hạn, có timeout và test khi đang thi."])
    replace_table(tables[9], risks)
    tests = [[c.text for c in row.cells] for row in tables[10].rows[1:]]
    tests[2] = ["TC03", "Sai mật khẩu/bị khóa; học sinh gửi CRUD/Start", "Từ chối auth/quyền; không thay đổi dữ liệu"]
    tests[3] = ["TC04", "Giảng viên CRUD rồi đăng nhập lại/restart Server", "Ngân hàng riêng phản ánh đúng dữ liệu đã lưu"]
    tests[5] = ["TC06", "Start thiếu Player ready hoặc thiếu câu", "Không bắt đầu; chủ phòng thấy nguyên nhân"]
    tests[12] = ["TC13", "Chủ phòng rớt mạng; không còn học sinh", "Ván tự chạy khi còn người; cùng tài khoản nhận lại điều khiển; hết người thì hủy"]
    tests.extend([
        ["TC17", "Chọn Giảng viên/Admin nhưng chưa login; tự gửi role", "Server không cấp quyền từ lựa chọn UI hoặc payload"],
        ["TC18", "A/B xem, tìm và sửa/xóa câu bằng ID tài khoản khác", "Mỗi người chỉ thấy/sửa câu mình; đoán ID không vượt quyền"],
        ["TC19", "Hai giảng viên tạo/điều khiển phòng và xem lịch sử", "Chỉ một phòng, đúng host; lịch sử lọc theo người tổ chức"],
        ["TC20", "Admin cấp/khóa/reset tài khoản; đổi mật khẩu tạm", "TEACHER không quản lý user; mật khẩu cũ/phiên bị thu hồi không dùng lại"],
        ["TC21", "Hai tài khoản cùng mật khẩu; kiểm tra DB và log", "Salt/hash khác nhau; không lưu/log mật khẩu rõ, không trả hash/salt qua API"],
        ["TC22", "Sai CA/SAN; thử kết nối không TLS; capture mạng", "Từ chối kết nối, không gửi login; không fallback plaintext; traffic hợp lệ mã hóa"],
        ["TC23", "Login sai liên tục khi ván đang chạy; Logout", "Rate limit/auth queue hữu hạn; game không bị chặn; quyền cũ được thu hồi"],
    ])
    replace_table(tables[10], tests)

    anchor = blocks[81]
    for text, is_heading in [
        ("5.9 Tài khoản và quyền sở hữu dữ liệu", True),
        ("Admin hệ thống: Tài khoản ADMIN đầu tiên được khởi tạo cục bộ trên Server, không có mật khẩu cố định trong mã nguồn. Admin tạo/khóa/reset tài khoản TEACHER; có thể tổ chức thi bằng ngân hàng của chính mình nhưng không mặc nhiên đọc/sửa câu hỏi của người khác.", False),
        ("Giảng viên: Đăng nhập bằng username/mật khẩu, quản lý ngân hàng và lịch sử của mình. Mỗi câu có owner_user_id, mỗi ván có host_user_id. Dữ liệu giữ trên cùng Server qua các lần đăng nhập; không đồng bộ sang Server khác.", False),
        ("Học sinh: Vai trò tham gia phòng, không có quyền câu hỏi hoặc tài khoản quản lý; phiên chỉ gắn nickname/player_id. Lựa chọn Bạn là ai chỉ dẫn tới màn hình phù hợp, không phải cơ chế xác thực.", False),
        ("Kiểm soát tại Server: Mọi truy vấn câu hỏi/lịch sử dùng user_id từ connection; sửa/xóa có điều kiện ID và owner. Chỉ chủ phòng được cấu hình, Start, Abort, Reset hoặc Close; giảng viên khác bị báo ROOM_BUSY/FORBIDDEN. Sau khi đóng phòng mới chuyển lượt tổ chức.", False),
        ("5.10 Bảo vệ mật khẩu và phiên đăng nhập", True),
        ("Khi lưu: Dùng PBKDF2WithHmacSHA256 có sẵn trong Java, 600.000 vòng ban đầu, salt SecureRandom tối thiểu 16 byte và hash 32 byte; lưu thuật toán/tham số cùng hash để nâng cấp. So sánh an toàn, không lưu mật khẩu rõ hoặc dùng MD5/SHA-256 đơn thuần [7].", False),
        ("Khi truyền: Dùng SSLServerSocket/SSLSocket với TLS 1.3 hoặc 1.2. Client truststore tin cậy CA của nhóm và kiểm tra SAN theo IP/hostname; private key chỉ ở Server. Không trust-all, không bỏ xác minh hoặc tự hạ xuống TCP không mã hóa khi lỗi [8].", False),
        ("Quản lý phiên: Server gắn user_id/role với Socket sau login; Logout/mất kết nối thu hồi quyền. Mật khẩu tạm bắt buộc đổi trước CRUD/tổ chức thi; đổi/reset/khóa thu hồi phiên cũ. Thao tác này với chủ phòng chỉ cho phép khi ván chưa chạy hoặc đã kết thúc. Giới hạn thử login và không ghi thông tin bí mật vào log.", False),
    ]:
        anchor = insert_after(anchor, blocks[77] if is_heading else blocks[71], text, label=not is_heading)
    anchor = blocks[119]
    for text in [
        "[7] OWASP Password Storage Cheat Sheet - https://cheatsheetseries.owasp.org/cheatsheets/Password_Storage_Cheat_Sheet.html",
        "[8] Java JSSE Reference Guide - https://docs.oracle.com/en/java/javase/21/security/java-secure-socket-extension-jsse-reference-guide.html",
    ]:
        anchor = insert_after(anchor, blocks[119], text)
    for drawing in blocks[37]._p.xpath('.//*[local-name()="docPr"]'):
        drawing.set("descr", "Client JavaFX giảng viên/Admin và học sinh kết nối Server Java qua TCP/TLS trong LAN; Server xác thực, kiểm tra owner và lưu tài khoản/câu hỏi/kết quả vào SQLite.")
    # Keep the scope heading/label with its first item, not isolated at page bottom.
    blocks[14].paragraph_format.keep_with_next = True
    doc.core_properties.modified = datetime.now(timezone.utc).replace(tzinfo=None)
    doc.core_properties.keywords = "Java, LAN, JavaFX, TCP, TLS, SQLite, tai khoan, phan quyen, ngan hang rieng"
    buf = BytesIO()
    doc.save(buf)
    with ZipFile(BytesIO(buf.getvalue())) as edited:
        replacement_parts = {
            "word/document.xml": edited.read("word/document.xml"),
            "docProps/core.xml": edited.read("docProps/core.xml"),
            "word/media/image1.png": diagram_bytes(),
        }
    output = BytesIO()
    with ZipFile(BytesIO(source_bytes)) as original, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for info in original.infolist():
            target.writestr(info, replacement_parts.get(info.filename, original_parts[info.filename]))
    result_bytes = output.getvalue()
    with ZipFile(BytesIO(result_bytes)) as result:
        changed = [name for name, data in original_parts.items() if result.read(name) != data]
        assert set(changed) == set(replacement_parts)
    assert sha256(REFERENCE.read_bytes()).hexdigest() == EXPECTED_HASH
    assert TARGET.read_bytes() in (source_bytes, result_bytes) or (TMP / "last_output.sha256").exists(), "Unexpected target changes"
    if (TMP / "last_output.sha256").exists():
        previous_hash = (TMP / "last_output.sha256").read_text(encoding="utf-8").strip()
        assert sha256(TARGET.read_bytes()).hexdigest() in (previous_hash, EXPECTED_HASH), "User modified target during QA"
    TARGET.write_bytes(result_bytes)
    final_hash = sha256(result_bytes).hexdigest()
    (TMP / "last_output.sha256").write_text(final_hash, encoding="utf-8")
    report = {
        "source_sha256": EXPECTED_HASH,
        "output_sha256": final_hash,
        "backup_unchanged": True,
        "changed_parts": changed,
        "preserved_parts": len(original_parts) - len(changed),
        "tables": len(doc.tables),
        "test_cases": len(tests),
        "application_code_changed": False,
    }
    (TMP / "structural_qa.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
