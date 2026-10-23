from pathlib import Path
import math
import textwrap

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"E:\_HUTECH\DO_AN_LAP_TRINH_MANG")
OUTPUT_DIR = ROOT / "output" / "documents"
ASSET_DIR = ROOT / "tmp" / "document_assets"
OUTPUT_PATH = OUTPUT_DIR / "De_cuong_Network_Quiz_Battle.docx"
DIAGRAM_PATH = ASSET_DIR / "kien_truc_network_quiz_battle.png"

NAVY = "1F4E78"
BLUE = "1F6FAB"
LIGHT_BLUE = "EAF3F9"
PALE_BLUE = "F5F9FC"
LIGHT_GRAY = "D9D9D9"
DARK_GRAY = "404040"
MID_GRAY = "737373"
WHITE = "FFFFFF"
BLACK = "000000"


def set_run_font(run, name="Arial", size=10.5, bold=None, color=BLACK, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_cm):
    cell.width = Cm(width_cm)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_GRAY, size=5):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_row_together(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def style_cell_text(cell, font_size=9.2, bold=False, color=BLACK, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)
    for paragraph in cell.paragraphs:
        paragraph.alignment = align
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.08
        for run in paragraph.runs:
            set_run_font(run, size=font_size, bold=bold, color=color)


def make_table(doc, headers, rows, widths, font_size=9.1, center_cols=None, header_fill=BLUE):
    center_cols = set(center_cols or [])
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    keep_row_together(hdr)
    for i, title in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = str(title)
        set_cell_width(cell, widths[i])
        set_cell_shading(cell, header_fill)
        style_cell_text(cell, font_size=font_size, bold=True, color=WHITE,
                        align=WD_ALIGN_PARAGRAPH.CENTER if i in center_cols else WD_ALIGN_PARAGRAPH.LEFT)
    for row_idx, values in enumerate(rows):
        row = table.add_row()
        keep_row_together(row)
        for col_idx, value in enumerate(values):
            cell = row.cells[col_idx]
            cell.text = str(value)
            set_cell_width(cell, widths[col_idx])
            if row_idx % 2 == 1:
                set_cell_shading(cell, LIGHT_BLUE)
            style_cell_text(
                cell,
                font_size=font_size,
                align=WD_ALIGN_PARAGRAPH.CENTER if col_idx in center_cols else WD_ALIGN_PARAGRAPH.LEFT,
            )
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_label_paragraph(doc, label, text, style=None, space_after=5):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(label)
    set_run_font(r, bold=True)
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_bullet(doc, text, level=0, font_size=10.5, space_after=3):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(text)
    set_run_font(r, size=font_size)
    return p


def add_numbered(doc, text, number, level=0, font_size=10.5, space_after=3):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.65 + 0.5 * level)
    p.paragraph_format.first_line_indent = Cm(-0.45)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(f"{number}.  {text}")
    set_run_font(r, size=font_size)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True
    r = p.add_run(text)
    set_run_font(r, size=15 if level == 1 else 12, bold=True, color=BLACK)
    return p


def add_field(paragraph, instruction):
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = instruction
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_run_font(run, size=8.5, color=MID_GRAY)


def add_footer(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(17.2), WD_TAB_ALIGNMENT.RIGHT)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), "4")
    top.set(qn("w:space"), "5")
    top.set(qn("w:color"), "B8CCE4")
    pBdr.append(top)
    pPr.append(pBdr)

    rr = p.add_run("Đề cương Network Quiz Battle - Môn Lập trình mạng máy tính")
    set_run_font(rr, size=8.5, color=MID_GRAY)
    rr = p.add_run("\tTrang ")
    set_run_font(rr, size=8.5, color=MID_GRAY)
    add_field(p, "PAGE")


def rounded_box(draw, xy, fill, outline, title, subtitle, fonts):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=24, fill=fill, outline=outline, width=4)
    title_font, body_font = fonts
    bbox = draw.textbbox((0, 0), title, font=title_font)
    tx = (x1 + x2 - (bbox[2] - bbox[0])) / 2
    draw.text((tx, y1 + 24), title, font=title_font, fill="#17365D")
    lines = []
    for para in subtitle.split("\n"):
        lines.extend(textwrap.wrap(para, width=25) or [""])
    y = y1 + 78
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=body_font)
        tx = (x1 + x2 - (bbox[2] - bbox[0])) / 2
        draw.text((tx, y), line, font=body_font, fill="#142536")
        y += 31


def draw_arrow(draw, start, end, color="#1F6FAB", width=8, both=False):
    draw.line([start, end], fill=color, width=width)
    def head(a, b):
        angle = math.atan2(b[1] - a[1], b[0] - a[0])
        length = 22
        wing = 0.55
        p1 = (b[0] - length * math.cos(angle - wing), b[1] - length * math.sin(angle - wing))
        p2 = (b[0] - length * math.cos(angle + wing), b[1] - length * math.sin(angle + wing))
        draw.polygon([b, p1, p2], fill=color)
    head(start, end)
    if both:
        head(end, start)


def create_architecture_diagram():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    img = Image.new("RGB", (1700, 700), "white")
    draw = ImageDraw.Draw(img)
    font_dir = Path(r"C:\Windows\Fonts")
    bold_path = font_dir / "arialbd.ttf"
    regular_path = font_dir / "arial.ttf"
    title_font = ImageFont.truetype(str(bold_path), 32)
    body_font = ImageFont.truetype(str(regular_path), 25)
    label_font = ImageFont.truetype(str(bold_path), 27)
    small_font = ImageFont.truetype(str(regular_path), 23)

    clients = [
        ((65, 70, 410, 225), "CLIENT 1", "Người chơi"),
        ((65, 275, 410, 430), "CLIENT 2", "Người chơi"),
        ((65, 480, 410, 635), "CLIENT N", "Người chơi / Quản trị viên"),
    ]
    for xy, title, subtitle in clients:
        rounded_box(draw, xy, "#DCEAF5", "#1F6FAB", title, subtitle, (label_font, small_font))

    rounded_box(draw, (600, 220, 940, 485), "#E8F2F8", "#1F6FAB", "MẠNG LAN",
                "Switch / Router\nTCP Port 5050\nJSON UTF-8", (label_font, small_font))
    rounded_box(draw, (1120, 85, 1615, 405), "#D8E8F4", "#17365D", "QUIZ SERVER",
                "Quản lý phiên\nPhòng thi\nQuiz Engine\nTính điểm và xếp hạng", (label_font, small_font))
    rounded_box(draw, (1180, 500, 1555, 655), "#EEF4F7", "#5B9BD5", "SQLITE",
                "Tài khoản, câu hỏi, kết quả", (label_font, small_font))

    for _, y in [(0, 147), (0, 352), (0, 557)]:
        draw_arrow(draw, (410, y), (600, 352), both=True)
    draw_arrow(draw, (940, 352), (1120, 248), both=True)
    draw_arrow(draw, (1368, 405), (1368, 500), color="#5B9BD5", both=True)
    draw.text((960, 405), "TCP hai chiều", font=small_font, fill="#17365D")
    draw.text((960, 442), "Server giữ trạng thái chuẩn", font=small_font, fill="#17365D")
    img.save(DIAGRAM_PATH, dpi=(220, 220))


def set_document_properties(doc):
    props = doc.core_properties
    props.title = "Đề cương chi tiết đề tài Network Quiz Battle"
    props.subject = "Đồ án môn học Lập trình mạng máy tính"
    props.author = "Nhóm sinh viên"
    props.keywords = "lập trình mạng, client-server, TCP, LAN, quiz, Java"
    props.comments = "Đề cương môn học"


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12

    title = styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for name, size, before, after in (("Heading 1", 15, 12, 6), ("Heading 2", 12, 8, 4)):
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Bullet 2", "List Number", "List Number 2"):
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(10.5)


def add_intro(doc):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_pr = p._p.get_or_add_pPr()
    direct_border = p_pr.find(qn("w:pBdr"))
    if direct_border is not None:
        p_pr.remove(direct_border)
    r = p.add_run("ĐỀ CƯƠNG CHI TIẾT ĐỀ TÀI")
    set_run_font(r, size=22, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("MÔN HỌC LẬP TRÌNH MẠNG MÁY TÍNH")
    set_run_font(r, size=15, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run("Đề tài: ")
    set_run_font(r, bold=True)
    r = p.add_run("Network Quiz Battle - Hệ thống thi trắc nghiệm nhiều người trong mạng LAN theo mô hình Client-Server, hỗ trợ quản lý ngân hàng câu hỏi, thi đấu thời gian thực và bảng xếp hạng")
    set_run_font(r, bold=True)

    add_label_paragraph(
        doc,
        "Mục đích: ",
        "Đề cương xác định phạm vi, chức năng, kiến trúc mạng, giao thức ứng dụng, công nghệ và kế hoạch triển khai để nhóm thống nhất cách thực hiện trước khi lập trình.",
    )
    add_label_paragraph(
        doc,
        "Kết quả chính: ",
        "Một Server điều phối phòng thi và nhiều Client trong cùng mạng LAN; câu hỏi, thời gian và bảng điểm được đồng bộ theo trạng thái do Server quản lý.",
        space_after=10,
    )


def build_document():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    create_architecture_diagram()

    doc = Document()
    set_document_properties(doc)
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.65)
    add_footer(section)

    add_intro(doc)

    add_heading(doc, "1 THÔNG TIN NHÓM")
    make_table(
        doc,
        ["Trường thông tin", "Nội dung"],
        [
            ["Tên nhóm", "................................................................................"],
            ["Lớp / học phần", "................................................................................"],
            ["Tên đề tài", "Network Quiz Battle - Hệ thống thi trắc nghiệm nhiều người trong mạng LAN"],
            ["Nhóm trưởng", "Họ tên - MSSV - nhiệm vụ"],
            ["Thành viên", "Họ tên - MSSV - nhiệm vụ của từng thành viên"],
            ["Thông tin liên hệ", "Email / số điện thoại nhóm trưởng"],
        ],
        [4.2, 13.0],
        font_size=9.5,
    )

    add_heading(doc, "2 GIỚI THIỆU ĐỀ TÀI")
    add_label_paragraph(
        doc,
        "Bối cảnh: ",
        "Trong giờ học hoặc sinh hoạt câu lạc bộ, việc tổ chức thi trắc nghiệm bằng giấy hoặc công cụ Internet khó bảo đảm hoạt động khi đường truyền ngoài không ổn định. Hệ thống trong LAN giúp giảng viên hoặc quản trị viên tổ chức một trận thi chung cho nhiều máy tính và quan sát kết quả ngay tại phòng máy.",
    )
    add_label_paragraph(
        doc,
        "Vấn đề cần giải quyết: ",
        "Nhiều Client phải nhận cùng nội dung và mốc thời gian, nộp đáp án trong giới hạn cho phép và xem bảng xếp hạng nhất quán. Server phải xử lý đồng thời nhiều kết nối, ngăn nộp trùng, chấm điểm tập trung và lưu kết quả.",
    )
    add_label_paragraph(
        doc,
        "Mục tiêu: ",
        "Xây dựng ứng dụng Java theo mô hình Client-Server dùng TCP, cho phép quản lý câu hỏi, tạo phòng, tham gia thi đấu thời gian thực, tính điểm theo độ chính xác và tốc độ, đồng thời lưu lịch sử trận đấu.",
    )
    add_label_paragraph(
        doc,
        "Đối tượng sử dụng: ",
        "Quản trị viên hoặc người tổ chức quản lý câu hỏi và phòng thi; người chơi đăng nhập, tham gia phòng và trả lời câu hỏi trên máy Client.",
    )
    add_label_paragraph(
        doc,
        "Sản phẩm dự kiến: ",
        "Ứng dụng Quiz Server, ứng dụng Quiz Client, cơ sở dữ liệu mẫu, mã nguồn, hướng dẫn cài đặt trong LAN, bộ dữ liệu kiểm thử, kịch bản demo và báo cáo đồ án.",
    )

    add_heading(doc, "3 PHẠM VI THỰC HIỆN")
    add_label_paragraph(doc, "Trong phạm vi: ", "")
    for item in [
        "Kết nối nhiều Client đến một Server trong cùng mạng LAN bằng TCP/IP.",
        "Đăng nhập cơ bản, phân biệt vai trò quản trị viên và người chơi.",
        "Thêm, sửa, xóa, tìm kiếm và nhập câu hỏi trắc nghiệm bốn lựa chọn từ tệp CSV theo mẫu.",
        "Tạo phòng thi, đặt tên phòng, bộ câu hỏi, số câu và thời gian trả lời cho từng câu.",
        "Tham gia phòng bằng mã phòng; hiển thị danh sách người chơi và trạng thái sẵn sàng.",
        "Phát câu hỏi, nhận đáp án, chấm điểm và cập nhật bảng xếp hạng gần thời gian thực.",
        "Lưu kết quả chung cuộc và tra cứu lịch sử trận đấu trên Server.",
        "Xử lý lỗi kết nối, timeout, thông điệp không hợp lệ và Client rời phòng.",
    ]:
        add_bullet(doc, item)
    add_label_paragraph(
        doc,
        "Ngoài phạm vi: ",
        "Thi qua Internet, ứng dụng di động, âm thanh hoặc video trực tiếp, chống gian lận ở mức giám sát màn hình, cân bằng tải nhiều Server, đồng bộ cloud và thanh toán. Các nội dung này chỉ thực hiện khi nhóm đăng ký mở rộng và được giảng viên duyệt.",
    )
    add_label_paragraph(
        doc,
        "Giả định và giới hạn: ",
        "Các máy ở cùng lớp mạng và cho phép truy cập cổng Server; mỗi người chơi dùng một tài khoản trên một Client tại một thời điểm; quy mô kiểm thử mục tiêu là 20 Client đồng thời.",
        space_after=2,
    )

    add_heading(doc, "4 DANH SÁCH CHỨC NĂNG HỆ THỐNG")
    p = doc.add_paragraph("Mỗi chức năng có mã riêng để thuận tiện phân công, kiểm thử và đối chiếu khi báo cáo.")
    p.paragraph_format.space_after = Pt(6)

    function_rows = [
        ["F01", "Kết nối Server", "Client nhập IP và Port, tạo kết nối TCP và thiết lập phiên.", "IP, Port / trạng thái", "Sai địa chỉ, timeout"],
        ["F02", "Đăng nhập và phân quyền", "Server xác thực tài khoản, trả vai trò và mã phiên.", "Tên, mật khẩu / phiên", "Sai thông tin, trùng phiên"],
        ["F03", "Quản lý câu hỏi", "Quản trị viên thêm, sửa, xóa, tìm kiếm hoặc nhập CSV.", "Nội dung câu hỏi / kết quả", "Thiếu đáp án, dữ liệu trùng"],
        ["F04", "Tạo và cấu hình phòng", "Chọn bộ câu hỏi, số câu, thời gian và quy tắc tính điểm.", "Cấu hình / mã phòng", "Cấu hình không hợp lệ"],
        ["F05", "Tham gia phòng", "Người chơi nhập mã phòng; Server kiểm tra và cập nhật sảnh chờ.", "Mã phòng / trạng thái sảnh", "Phòng đầy, đang thi"],
        ["F06", "Sẵn sàng và bắt đầu", "Server kiểm tra điều kiện rồi phát COUNTDOWN và QUIZ_START.", "Trạng thái / mốc bắt đầu", "Chưa đủ người, mất chủ phòng"],
        ["F07", "Phát câu hỏi đồng bộ", "Server gửi cùng câu hỏi, lựa chọn và mốc thời gian đến các Client.", "Câu hỏi / QUESTION", "Client chậm, mất kết nối"],
        ["F08", "Nhận và khóa đáp án", "Server nhận đáp án đầu tiên hợp lệ, xác nhận và từ chối nộp trùng.", "Đáp án / ANSWER_ACK", "Quá hạn, trùng request"],
        ["F09", "Chấm điểm", "Server kiểm tra đáp án và tính điểm theo thời gian nhận tại Server.", "Đáp án, thời gian / điểm", "Dữ liệu câu hỏi lỗi"],
        ["F10", "Bảng xếp hạng trực tiếp", "Sau mỗi câu, Server sắp hạng và phát LEADERBOARD.", "Kết quả câu / thứ hạng", "Đồng điểm, người chơi rời"],
        ["F11", "Kết thúc và lưu lịch sử", "Tổng hợp kết quả, lưu trận đấu và hiển thị xếp hạng cuối.", "Phiên thi / kết quả lưu", "Lỗi cơ sở dữ liệu"],
        ["F12", "Giám sát và xử lý lỗi", "Heartbeat, đóng tài nguyên, ghi nhật ký và phản hồi mã lỗi.", "Sự kiện / ERROR hoặc phục hồi", "Ngắt đột ngột, gói sai"],
    ]
    make_table(
        doc,
        ["Mã", "Chức năng", "Xử lý chính", "Đầu vào / đầu ra", "Lỗi chính"],
        function_rows,
        [1.1, 3.1, 6.4, 3.4, 3.2],
        font_size=8.25,
        center_cols={0},
    )

    add_heading(doc, "4.1 Đặc tả chức năng trọng tâm", level=2)
    add_label_paragraph(doc, "F07 Phát câu hỏi đồng bộ: ", "Server chọn câu hỏi theo cấu hình phòng, tạo mã vòng, ghi thời điểm bắt đầu và hạn trả lời theo đồng hồ Server, sau đó gửi thông điệp QUESTION cho mọi người chơi đang hoạt động. Thông điệp chỉ chứa nội dung và các lựa chọn, không chứa đáp án đúng.")
    add_label_paragraph(doc, "F08 Nhận và khóa đáp án: ", "Client gửi roomId, questionId, lựa chọn và requestId. Server kiểm tra trạng thái phòng, hạn trả lời và quyền nộp; đáp án hợp lệ đầu tiên được ghi nhận. requestId giúp Server trả lại cùng kết quả khi Client gửi lại do chưa nhận được xác nhận.")
    add_label_paragraph(doc, "F09 Chấm điểm: ", "Đáp án sai hoặc đến sau hạn nhận 0 điểm. Đáp án đúng nhận 1.000 điểm cơ bản và tối đa 500 điểm tốc độ. Điểm tốc độ giảm tuyến tính theo thời gian phản hồi do Server đo. Công thức có thể điều chỉnh trước khi giảng viên duyệt nhưng phải cố định trong suốt một trận.")
    add_label_paragraph(doc, "F10 Xếp hạng: ", "Thứ hạng ưu tiên tổng điểm, sau đó đến số câu đúng và tổng thời gian trả lời. Nếu vẫn bằng nhau, hai người chơi được xem là đồng hạng. Bảng xếp hạng sau mỗi câu và kết quả cuối đều do Server phát để mọi Client hiển thị cùng dữ liệu.")

    add_heading(doc, "4.2 Tiêu chí chấp nhận chức năng trọng tâm", level=2)
    acceptance_items = [
        "Ba Client nhận cùng questionId và cùng thời hạn; chênh lệch hiển thị không làm thay đổi thời điểm khóa tại Server.",
        "Một Client gửi hai ANSWER_SUBMIT cho cùng câu chỉ được ghi nhận một đáp án.",
        "Client nộp sau hạn nhận được lỗi ANSWER_TOO_LATE và không tăng điểm.",
        "Bảng xếp hạng trên các Client khớp với dữ liệu Server sau mỗi câu và khi kết thúc.",
    ]
    p = doc.add_paragraph("; ".join(
        item.rstrip(".") if idx == 0 else item[0].lower() + item[1:].rstrip(".")
        for idx, item in enumerate(acceptance_items)
    ) + ".")
    p.paragraph_format.space_after = Pt(4)
    for r in p.runs:
        set_run_font(r, size=10.2)

    add_heading(doc, "5 KIẾN TRÚC VÀ LUỒNG DỮ LIỆU")
    add_heading(doc, "5.1 Mô hình tổng thể", level=2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(DIAGRAM_PATH), width=Cm(17.0))
    doc_pr = run._r.xpath(".//wp:docPr")
    if doc_pr:
        doc_pr[0].set("descr", "Sơ đồ nhiều ứng dụng Client kết nối Quiz Server trong mạng LAN; Server truy cập cơ sở dữ liệu SQLite")
    p.paragraph_format.space_after = Pt(3)
    cap = doc.add_paragraph("Hình 1 Mô hình triển khai Network Quiz Battle trong mạng LAN")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    for r in cap.runs:
        set_run_font(r, size=9, italic=True, color=DARK_GRAY)

    add_label_paragraph(
        doc,
        "Nguyên tắc: ",
        "Server là nguồn dữ liệu chuẩn cho trạng thái phòng, thời gian, đáp án đúng, điểm và thứ hạng. Client chỉ gửi yêu cầu và hiển thị phản hồi; Client không tự quyết định điểm hoặc thời điểm kết thúc câu hỏi.",
    )

    component_rows = [
        ["Ứng dụng Client", "Giao diện đăng nhập, sảnh chờ, câu hỏi, đồng hồ đếm ngược và bảng điểm; NetworkClient đóng gói và nhận thông điệp."],
        ["Connection Acceptor", "Lắng nghe TCP Port, chấp nhận Socket và giao phiên kết nối cho ExecutorService."],
        ["Session Manager", "Quản lý tài khoản đang đăng nhập, heartbeat, requestId và trạng thái kết nối."],
        ["Room Manager", "Tạo phòng, thêm hoặc loại người chơi, quản lý chủ phòng và trạng thái sảnh."],
        ["Quiz Engine", "Chọn câu hỏi, lập lịch mở hoặc đóng câu, kiểm tra đáp án và chuyển trạng thái trận."],
        ["Scoring Service", "Tính điểm tập trung, xử lý đồng điểm và tạo bảng xếp hạng."],
        ["Data Repository", "Đọc và ghi tài khoản, câu hỏi, phòng, trận đấu và kết quả vào SQLite."],
        ["Logger", "Ghi sự kiện kết nối, lỗi giao thức và lỗi dữ liệu để phục vụ kiểm thử."],
    ]
    make_table(doc, ["Thành phần", "Trách nhiệm"], component_rows, [4.3, 12.9], font_size=9.0)

    add_heading(doc, "5.2 Trạng thái phòng và vòng thi", level=2)
    p = doc.add_paragraph("WAITING -> COUNTDOWN -> QUESTION_OPEN -> QUESTION_RESULT -> QUESTION_OPEN hoặc FINISHED")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    for r in p.runs:
        set_run_font(r, name="Consolas", size=9.5, bold=True)
    add_label_paragraph(doc, "Quy tắc chuyển trạng thái: ", "Chỉ Quiz Engine trên Server được chuyển trạng thái. ANSWER_SUBMIT chỉ hợp lệ trong QUESTION_OPEN. Khi hết hạn hoặc tất cả người chơi hoạt động đã trả lời, Server chuyển sang QUESTION_RESULT, phát đáp án đúng và bảng xếp hạng rồi mở câu tiếp theo.")

    add_heading(doc, "5.3 Giao thức ứng dụng", level=2)
    add_label_paragraph(doc, "Kênh truyền: ", "TCP, cổng đề xuất 5050 và cho phép thay đổi trong tệp cấu hình Server. Mỗi thông điệp gồm 4 byte độ dài theo thứ tự byte mạng, tiếp theo là chuỗi JSON UTF-8 có đúng số byte đã khai báo. Server từ chối khung vượt quá 256 KB.")
    add_label_paragraph(doc, "Cấu trúc chung: ", "type xác định loại thông điệp; requestId dùng đối chiếu yêu cầu và chống xử lý lặp; sessionId xác định phiên đăng nhập; roomId xác định phòng; payload chứa dữ liệu nghiệp vụ; sentAt là thời điểm tham khảo của bên gửi.")
    code_lines = [
        '{',
        '  "type": "ANSWER_SUBMIT",',
        '  "requestId": "req-1042",',
        '  "sessionId": "s-8f31",',
        '  "roomId": "QZ7K2",',
        '  "payload": {"questionId": 25, "choice": "B"},',
        '  "sentAt": 1791421205123',
        '}',
    ]
    for line in code_lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1.0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(line)
        set_run_font(r, name="Consolas", size=8.8)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

    message_rows = [
        ["Kết nối", "AUTH_REQUEST, AUTH_RESULT, PING, PONG, LOGOUT"],
        ["Phòng", "ROOM_CREATE, ROOM_JOIN, ROOM_LEAVE, LOBBY_UPDATE, READY_SET"],
        ["Thi đấu", "QUIZ_START, QUESTION, ANSWER_SUBMIT, ANSWER_ACK, QUESTION_RESULT"],
        ["Kết quả", "LEADERBOARD, QUIZ_FINISHED, MATCH_HISTORY"],
        ["Lỗi", "ERROR với code, message và requestId liên quan"],
    ]
    make_table(doc, ["Nhóm", "Loại thông điệp chính"], message_rows, [3.4, 13.8], font_size=9.0)

    add_heading(doc, "5.4 Luồng thi đấu chính", level=2)
    for idx, item in enumerate([
        "Client kết nối TCP, gửi AUTH_REQUEST và nhận sessionId nếu xác thực thành công.",
        "Chủ phòng gửi ROOM_CREATE; người chơi gửi ROOM_JOIN và READY_SET.",
        "Server phát QUIZ_START và mốc đếm ngược; mọi Client chuyển sang màn hình thi.",
        "Với từng câu, Server phát QUESTION, nhận ANSWER_SUBMIT và trả ANSWER_ACK.",
        "Hết thời gian, Server chấm điểm, phát QUESTION_RESULT và LEADERBOARD.",
        "Sau câu cuối, Server phát QUIZ_FINISHED, lưu kết quả rồi cho phép rời hoặc chơi lại.",
    ], start=1):
        add_numbered(doc, item, idx)

    add_heading(doc, "5.5 Xử lý đồng thời và lỗi kết nối", level=2)
    for item in [
        "Server dùng ExecutorService để xử lý nhiều Socket; luồng giao diện Client không thực hiện I/O mạng trực tiếp.",
        "Mỗi phòng có khóa hoặc hàng đợi sự kiện riêng để các thao tác bắt đầu, nộp đáp án và kết thúc câu được xử lý tuần tự.",
        "ScheduledExecutorService quản lý countdown, hạn câu hỏi và heartbeat; thời gian Server là chuẩn chấm điểm.",
        "Nếu mất PONG quá số chu kỳ cấu hình, Server đánh dấu Client ngắt kết nối, giải phóng tài nguyên và cập nhật sảnh hoặc bảng điểm.",
        "Thông điệp sai JSON, sai trạng thái hoặc không đủ quyền nhận ERROR; lỗi của một Client không được làm dừng tiến trình Server.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "5.6 Dữ liệu lưu trữ", level=2)
    data_rows = [
        ["users", "id, username, password_hash, display_name, role, status"],
        ["questions", "id, content, option_a..d, correct_option, difficulty, category, created_by"],
        ["rooms", "id, code, host_user_id, state, question_count, time_limit_sec"],
        ["matches", "id, room_id, started_at, ended_at, total_questions"],
        ["match_answers", "match_id, question_id, user_id, selected_option, is_correct, response_ms, score"],
        ["match_results", "match_id, user_id, total_score, correct_count, total_response_ms, rank"],
    ]
    make_table(doc, ["Bảng", "Dữ liệu chính"], data_rows, [4.0, 13.2], font_size=8.8)

    add_heading(doc, "6 CÔNG NGHỆ SỬ DỤNG")
    tech_rows = [
        ["Ngôn ngữ", "Java 17", "Phát triển Client và Server; phù hợp nội dung Socket, Thread và Stream của học phần."],
        ["Giao diện", "Java Swing", "Tạo giao diện desktop, chạy với JDK mà không cần cài thêm bộ JavaFX."],
        ["Kết nối", "Socket, ServerSocket, TCP/IP", "Thiết lập kênh tin cậy trong LAN và minh họa lập trình mạng mức ứng dụng."],
        ["Định dạng", "JSON UTF-8, khung độ dài 4 byte", "Trao đổi thông điệp có cấu trúc và xử lý đúng ranh giới dữ liệu trên TCP."],
        ["Đồng thời", "ExecutorService, ScheduledExecutorService", "Phục vụ nhiều Client và điều phối thời gian câu hỏi."],
        ["Đồng bộ dữ liệu", "ConcurrentHashMap, khóa theo phòng", "Giữ trạng thái phiên và xử lý nộp đáp án an toàn."],
        ["Cơ sở dữ liệu", "SQLite, JDBC", "Lưu tài khoản, ngân hàng câu hỏi và lịch sử thi bằng một tệp dữ liệu."],
        ["Thư viện", "Gson hoặc Jackson", "Tuần tự hóa và phân tích JSON; nhóm chọn một thư viện duy nhất."],
        ["Xây dựng", "Maven", "Quản lý thư viện, cấu trúc dự án và đóng gói ứng dụng."],
        ["Kiểm thử", "JUnit 5, Wireshark", "Kiểm thử logic chấm điểm và quan sát luồng TCP khi chạy trong LAN."],
        ["Công cụ", "NetBeans hoặc IntelliJ IDEA, Git", "Lập trình, quản lý phiên bản và tích hợp mã nguồn."],
    ]
    make_table(doc, ["Thành phần", "Lựa chọn đề xuất", "Lý do và cách sử dụng"], tech_rows, [3.2, 4.6, 9.4], font_size=8.8)

    add_heading(doc, "7 LIÊN HỆ KIẾN THỨC MÔN HỌC")
    course_rows = [
        ["F01, F02", "TCP/IP, Socket, ServerSocket", "Giải thích bắt tay ứng dụng, vòng đời Socket và phiên đăng nhập."],
        ["F07, F08", "InputStream, OutputStream, DataInputStream, DataOutputStream", "Chỉ ra phần đọc 4 byte độ dài và khung JSON UTF-8."],
        ["F01 - F12", "Giao thức ứng dụng", "Trình bày message type, requestId, trạng thái và mã lỗi."],
        ["F05 - F10", "Thread, ExecutorService, đồng bộ", "Demo nhiều Client vào cùng phòng và nộp đáp án đồng thời."],
        ["F06 - F10", "Lập lịch và timeout", "Giải thích countdown, hạn trả lời, heartbeat và đồng hồ Server."],
        ["F03, F11", "JDBC và lưu trữ dữ liệu", "Truy vấn câu hỏi, transaction lưu kết quả và xử lý lỗi cơ sở dữ liệu."],
        ["F12", "Exception, đóng tài nguyên", "Demo mất kết nối, gói sai và Server tiếp tục hoạt động."],
        ["Toàn hệ thống", "Phân tích gói tin", "Dùng Wireshark xác minh địa chỉ IP, Port và chuỗi gói TCP."],
    ]
    make_table(doc, ["Chức năng", "Kiến thức vận dụng", "Minh chứng khi báo cáo"], course_rows, [3.0, 5.5, 8.7], font_size=8.8)

    add_heading(doc, "8 KẾ HOẠCH TRIỂN KHAI VÀ PHÂN CÔNG")
    plan_rows = [
        ["Tuần 1", "Phân tích yêu cầu, thống nhất phạm vi, hoàn thiện đề cương", "Đề cương được giảng viên góp ý", "Cả nhóm"],
        ["Tuần 2", "Thiết kế giao thức, kiến trúc, cơ sở dữ liệu và giao diện", "Sơ đồ, message catalog, lược đồ DB", "Nhóm trưởng và thành viên"],
        ["Tuần 3", "Xây dựng Server kết nối, phiên đăng nhập và quản lý phòng", "F01, F02, F04, F05", "Phân công theo bảng dưới"],
        ["Tuần 4", "Xây dựng Client, ngân hàng câu hỏi và sảnh chờ", "F03, giao diện F05, F06", "Phân công theo bảng dưới"],
        ["Tuần 5", "Hoàn thiện Quiz Engine, tính điểm và bảng xếp hạng", "F07 - F11", "Phân công theo bảng dưới"],
        ["Tuần 6", "Tích hợp, kiểm thử nhiều Client, xử lý mất kết nối", "Bản tích hợp và biên bản lỗi", "Cả nhóm"],
        ["Tuần 7", "Tối ưu giao diện, hoàn thiện dữ liệu mẫu và tài liệu", "Bản release candidate", "Cả nhóm"],
        ["Tuần 8", "Tổng duyệt, quay video, hoàn thiện báo cáo và nộp", "Bộ sản phẩm cuối", "Cả nhóm"],
    ]
    make_table(doc, ["Giai đoạn", "Công việc", "Sản phẩm", "Người phụ trách"], plan_rows, [2.1, 7.1, 4.7, 3.3], font_size=8.45, center_cols={0})

    add_heading(doc, "8.1 Bảng phân công thành viên", level=2)
    assignment_rows = [
        ["Nhóm trưởng", "........................", "........................", "F01, F02, F04", "Server phiên và phòng"],
        ["Thành viên 2", "........................", "........................", "F03", "Ngân hàng câu hỏi và DB"],
        ["Thành viên 3", "........................", "........................", "F05 - F08", "Client, sảnh và thi đấu"],
        ["Thành viên 4", "........................", "........................", "F09 - F12", "Điểm, xếp hạng, kiểm thử"],
    ]
    make_table(doc, ["Vai trò", "Họ tên", "MSSV", "Mã chức năng", "Sản phẩm bàn giao"], assignment_rows, [2.8, 3.5, 2.8, 3.1, 5.0], font_size=8.6, center_cols={2, 3})

    add_heading(doc, "8.2 Rủi ro và phương án xử lý", level=2)
    risk_rows = [
        ["Giao thức thay đổi muộn", "Cao", "Chốt message catalog ở tuần 2; lưu phiên bản giao thức trong trường protocolVersion."],
        ["Xung đột xử lý đồng thời", "Cao", "Khóa theo phòng, requestId idempotent và kiểm thử nhiều Client gửi cùng lúc."],
        ["Chênh lệch đồng hồ Client", "Trung bình", "Server gửi mốc thời gian và quyết định hạn; Client chỉ hiển thị đếm ngược."],
        ["Mất kết nối khi đang thi", "Trung bình", "Heartbeat, timeout, đóng tài nguyên và giữ kết quả đã được Server xác nhận."],
        ["Dữ liệu câu hỏi sai", "Trung bình", "Ràng buộc trường bắt buộc, kiểm tra đáp án đúng và transaction khi nhập CSV."],
        ["Trễ tiến độ", "Trung bình", "Chia chức năng theo mã, tích hợp hằng tuần và ưu tiên phạm vi cốt lõi F01 - F12."],
    ]
    make_table(doc, ["Rủi ro", "Mức độ", "Phương án xử lý"], risk_rows, [4.6, 2.4, 10.2], font_size=8.7, center_cols={1})

    add_heading(doc, "9 KỊCH BẢN KIỂM THỬ VÀ DEMO")
    test_rows = [
        ["TC01", "Client nhập đúng IP và Port", "Kết nối thành công, nhận AUTH_RESULT"],
        ["TC02", "Sai IP, Port hoặc Server chưa chạy", "Báo lỗi rõ ràng, giao diện không treo"],
        ["TC03", "Đăng nhập sai hoặc đăng nhập trùng", "Server từ chối bằng mã lỗi phù hợp"],
        ["TC04", "Quản trị viên thêm và sửa câu hỏi", "Dữ liệu hợp lệ được lưu và đọc lại đúng"],
        ["TC05", "Nhập CSV có dòng sai cấu trúc", "Báo dòng lỗi; không lưu dữ liệu hỏng"],
        ["TC06", "Tạo phòng và ba Client tham gia", "Danh sách sảnh giống nhau trên các Client"],
        ["TC07", "Tham gia sai mã hoặc phòng đã bắt đầu", "Server từ chối, Client vẫn hoạt động"],
        ["TC08", "Bắt đầu khi chưa đủ điều kiện", "Server không chuyển trạng thái và nêu lý do"],
        ["TC09", "Ba Client nhận câu hỏi", "Cùng questionId, nội dung và hạn trả lời"],
        ["TC10", "Hai Client nộp gần như đồng thời", "Mỗi đáp án được xử lý độc lập, không sai điểm"],
        ["TC11", "Một Client nộp hai lần", "Chỉ đáp án đầu hợp lệ được ghi nhận"],
        ["TC12", "Nộp sau khi hết giờ", "Nhận ANSWER_TOO_LATE, điểm không thay đổi"],
        ["TC13", "Ngắt một Client giữa trận", "Server tiếp tục phục vụ Client còn lại"],
        ["TC14", "Gửi JSON sai hoặc frame quá lớn", "Server trả lỗi hoặc đóng phiên, không dừng toàn hệ thống"],
        ["TC15", "Kết thúc trận", "Thứ hạng các Client khớp DB và quy tắc đồng điểm"],
        ["TC16", "20 Client mô phỏng hoạt động", "Server ổn định, không mất đáp án đã xác nhận"],
    ]
    make_table(doc, ["Mã test", "Tình huống", "Kết quả mong đợi"], test_rows, [2.0, 7.2, 8.0], font_size=8.7, center_cols={0})

    add_heading(doc, "9.1 Kịch bản demo đề xuất", level=2)
    for idx, item in enumerate([
        "Khởi động Server, hiển thị IP LAN và Port đang lắng nghe.",
        "Quản trị viên đăng nhập, thêm một câu hỏi và tạo phòng thi năm câu.",
        "Ít nhất ba Client từ các máy khác nhau đăng nhập, nhập mã phòng và sẵn sàng.",
        "Bắt đầu trận, trả lời đúng hoặc sai với tốc độ khác nhau để quan sát cách tính điểm.",
        "Ngắt một Client ở giữa trận, xác minh Server và các Client còn lại tiếp tục hoạt động.",
        "Kết thúc trận, đối chiếu bảng xếp hạng trên Client với lịch sử lưu tại Server.",
        "Mở Wireshark để chỉ ra kết nối TCP, IP Server, Port và các gói dữ liệu trao đổi.",
    ], start=1):
        add_numbered(doc, item, idx)

    add_heading(doc, "10 TIÊU CHÍ ĐÁNH GIÁ THAM KHẢO")
    rubric_rows = [
        ["Chức năng đã đăng ký hoạt động đúng", "35%"],
        ["Vận dụng kiến thức lập trình mạng", "30%"],
        ["Kiến trúc hệ thống và chất lượng mã nguồn", "15%"],
        ["Demo, kiểm thử và xử lý lỗi", "10%"],
        ["Giải thích và trả lời câu hỏi", "10%"],
        ["Tổng cộng", "100%"],
    ]
    table = make_table(doc, ["Tiêu chí", "Tỷ trọng"], rubric_rows, [14.2, 3.0], font_size=9.3, center_cols={1})
    for cell in table.rows[-1].cells:
        set_cell_shading(cell, "D9EAF7")
        for r in cell.paragraphs[0].runs:
            r.bold = True

    add_heading(doc, "11 QUY ĐỊNH NỘP ĐỀ CƯƠNG VÀ BÁO CÁO")
    for item in [
        "Nộp đề cương trước khi triển khai; mọi thay đổi lớn về phạm vi hoặc chức năng cần được giảng viên đồng ý.",
        "Mã nguồn phải tách Client, Server và lớp dùng chung; thông tin IP, Port, đường dẫn cơ sở dữ liệu phải cấu hình được.",
        "Mỗi chức năng F01 - F12 phải có người phụ trách, tiêu chí hoàn thành và minh chứng demo.",
        "Nhóm nộp mã nguồn, tệp cấu hình, cơ sở dữ liệu mẫu, hướng dẫn chạy, báo cáo, kịch bản kiểm thử và video nếu giảng viên yêu cầu.",
        "Tài khoản hoặc dữ liệu mẫu không chứa thông tin nhạy cảm; mật khẩu lưu trong cơ sở dữ liệu phải ở dạng băm.",
        "Được phép sử dụng công cụ AI hỗ trợ nhưng các thành viên phải giải thích được mã nguồn, giao thức, luồng dữ liệu, xử lý đồng thời và lỗi kết nối.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "12 PHẦN XÁC NHẬN DUYỆT ĐỀ CƯƠNG")
    approval_rows = [
        ["Đánh giá phạm vi và tính phù hợp", "................................................................................\n................................................................................"],
        ["Các chức năng cần điều chỉnh", "................................................................................\n................................................................................"],
        ["Trạng thái", "[  ] Duyệt     [  ] Duyệt có chỉnh sửa     [  ] Chưa duyệt"],
        ["Ngày duyệt / Giảng viên", "................................................................................"],
    ]
    make_table(doc, ["Ý kiến giảng viên", "Kết quả"], approval_rows, [7.0, 10.2], font_size=9.3)

    # Encourage Word/LibreOffice to refresh fields such as page numbers.
    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build_document()
