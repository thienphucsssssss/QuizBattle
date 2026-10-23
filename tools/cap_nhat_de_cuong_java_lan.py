"""Edit the retained proposal; preserve all non-editable OOXML package parts."""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import json
import math

from docx import Document
from docx.text.paragraph import Paragraph
from docx.shared import RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "output/documents/De_cuong_Network_Quiz_Battle_Nop_Giang_Vien.docx"
TARGET = ROOT / "output/documents/De_cuong_Network_Quiz_Battle_Java_LAN_Nop_Giang_Vien.docx"
TMP = ROOT / "tmp/de_cuong_java_lan"
EXPECTED_HASH = "502fd30f801a63aef4fc4d8e0152e27e9c16569b5c55084d181f8e485bbda16b"


def replace_paragraph(p, text, label=False):
    rpr = deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    p.clear()
    parts = text.split(": ", 1) if label and ": " in text else [text]
    for i, part in enumerate(parts):
        run = p.add_run(part + (": " if len(parts) == 2 and i == 0 else ""))
        if rpr is not None:
            run._r.insert(0, deepcopy(rpr))
        run.font.color.rgb = RGBColor(0, 0, 0)
        if len(parts) == 2:
            run.bold = i == 0


def replace_cell(cell, text):
    first = cell.paragraphs[0]
    rpr = deepcopy(first.runs[0]._r.rPr) if first.runs and first.runs[0]._r.rPr is not None else None
    for p in cell.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    first.clear()
    run = first.add_run(text)
    if rpr is not None:
        run._r.insert(0, rpr)


def replace_table(table, records):
    templates = [deepcopy(row._tr) for row in table.rows[1:3]]
    while len(table.rows) > len(records) + 1:
        table._tbl.remove(table.rows[-1]._tr)
    while len(table.rows) < len(records) + 1:
        table._tbl.append(deepcopy(templates[(len(table.rows) - 1) % len(templates)]))
    for row, record in zip(table.rows[1:], records):
        for cell, value in zip(row.cells, record):
            replace_cell(cell, value)


def diagram_bytes():
    img = Image.new("RGB", (1700, 700), "white")
    draw = ImageDraw.Draw(img)
    fonts = Path(r"C:\Windows\Fonts")
    bold = ImageFont.truetype(str(fonts / "arialbd.ttf"), 30)
    body = ImageFont.truetype(str(fonts / "arial.ttf"), 25)
    small = ImageFont.truetype(str(fonts / "arial.ttf"), 23)

    def box(xy, title, lines):
        x1, y1, x2, y2 = xy
        draw.rounded_rectangle(xy, radius=20, fill="#EAF3F9", outline="#1F6FAB", width=4)
        for i, (line, font) in enumerate([(title, bold)] + [(line, body) for line in lines]):
            width = draw.textlength(line, font=font)
            draw.text(((x1 + x2 - width) / 2, y1 + 23 + i * 37), line, font=font, fill="#17365D")

    def arrow(a, b):
        draw.line([a, b], fill="#1F6FAB", width=7)
        for start, end in [(a, b), (b, a)]:
            angle = math.atan2(end[1] - start[1], end[0] - start[0])
            points = [end] + [(end[0] - 19 * math.cos(angle + d), end[1] - 19 * math.sin(angle + d)) for d in [-0.52, 0.52]]
            draw.polygon(points, fill="#1F6FAB")

    box((55, 35, 450, 180), "ADMIN APP", ["Java và JavaFX", "Quản trị qua TCP Socket"])
    box((55, 220, 450, 365), "PLAYER APP 1", ["Java và JavaFX", "Tham gia và trả lời"])
    box((55, 405, 450, 550), "PLAYER APP 2 ĐẾN 10", ["Cùng phòng thi", "Hiển thị điểm từ Server"])
    box((610, 160, 975, 420), "MẠNG LAN", ["Switch hoặc Wi-Fi", "TCP Socket hai chiều", "JSON UTF-8 từng dòng", "Cổng mặc định 5050"])
    box((1130, 55, 1660, 335), "JAVA SERVER", ["Một máy trong LAN", "ServerSocket và worker", "Game Worker và deadline", "Chấm điểm và xếp hạng"])
    box((1165, 435, 1625, 580), "SQLITE QUA JDBC", ["Câu hỏi và kết quả", "File tại máy Server"])
    for y in [107, 292, 477]:
        arrow((450, y), (610, 290))
    arrow((975, 290), (1130, 197))
    arrow((1395, 335), (1395, 435))
    draw.text((65, 605), "Ví dụ: máy Server 192.168.1.10, các Client cùng LAN kết nối tới 192.168.1.10:5050.", font=small, fill="#17365D")
    draw.text((65, 645), "Admin có thể chạy cùng máy Server; chỉ Server được đọc ghi SQLite. Không cần Internet.", font=small, fill="#17365D")
    buf = BytesIO()
    img.save(buf, format="PNG", dpi=(220, 220))
    return buf.getvalue()


def main():
    source_bytes = SOURCE.read_bytes()
    assert sha256(source_bytes).hexdigest() == EXPECTED_HASH, "Source changed; inspect new user edits first"
    TMP.mkdir(parents=True, exist_ok=True)
    with ZipFile(BytesIO(source_bytes)) as z:
        manifest = {p.filename: {"size": p.file_size, "sha256": sha256(z.read(p.filename)).hexdigest()} for p in z.infolist()}
    (TMP / "manifest.json").write_text(json.dumps({"source_sha256": EXPECTED_HASH, "parts": manifest}, indent=2), encoding="utf-8")
    doc = Document(BytesIO(source_bytes))
    blocks = list(doc.iter_inner_content())
    tables = list(doc.tables)

    updates = {
        2: "Đề tài: Network Quiz Battle - Hệ thống thi trắc nghiệm nhiều người trong mạng LAN theo mô hình Client-Server, hỗ trợ quản lý ngân hàng câu hỏi, thi đấu thời gian thực và bảng xếp hạng; toàn bộ ứng dụng viết bằng Java.",
        3: "Mục đích: Nhóm đăng ký xây dựng phần mềm desktop bằng Java, xác định chức năng, kiến trúc mạng và kế hoạch triển khai để giảng viên xem xét trước khi lập trình.",
        4: "Kết quả chính: Server Java chạy trên một máy trong LAN; Admin và Player dùng ứng dụng JavaFX kết nối bằng TCP Socket. SQLite lưu ngân hàng câu hỏi và kết quả tại máy Server.",
        8: "Bối cảnh: Trong lớp học hoặc sinh hoạt nhóm, người tổ chức cần quản lý câu hỏi và điều hành một ván trắc nghiệm chung. Nhóm xây dựng ứng dụng dùng trong mạng LAN, hoạt động không cần kết nối Internet.",
        9: "Vấn đề cần giải quyết: Server phải phục vụ nhiều Client, phát cùng câu hỏi, nhận đáp án đúng hạn, ngăn nộp lặp và trả bảng điểm nhất quán. Người tổ chức điều khiển ván thi bằng Admin App qua cùng kênh TCP.",
        10: "Mục tiêu: Sử dụng Java và TCP Socket để vận dụng kiến thức lập trình mạng; xây giao diện desktop bằng JavaFX, quản lý câu hỏi qua SQLite JDBC, tổ chức ván thi và lưu kết quả chung cuộc. Server quyết định thời hạn, điểm và thứ hạng.",
        12: "Sản phẩm dự kiến: Server console, Admin App, Player App, ít nhất 20 câu hỏi mẫu, hướng dẫn chạy trên các máy trong LAN, bộ kiểm thử, video demo và báo cáo đồ án.",
        15: "Một Server Java chạy trên máy trong LAN; Admin App và Player App viết bằng JavaFX, kết nối bằng TCP Socket.",
        16: "Xác thực một Admin điều hành trong LAN; Player tham gia bằng nickname và mã phòng. Chỉ Server được truy cập cơ sở dữ liệu.",
        22: "Xử lý sai dữ liệu, lỗi kết nối, Client rời phòng và ghi log; kiểm thử Socket trên nhiều máy trong LAN.",
        23: "Ngoài phạm vi: Triển khai Server qua Internet, nhiều phòng đồng thời, tài khoản Player, khôi phục phiên thi, web/mobile, chat/voice, lịch sử từng đáp án và chống gian lận nâng cao. Nhập CSV và bộ cài tự động là hướng mở rộng.",
        24: "Giả định và giới hạn: Một Admin, tối thiểu 3 Player khi nghiệm thu; giới hạn ban đầu 10 Player để kiểm thử. Các máy cùng LAN và truy cập được cổng TCP của Server. Máy Server phải duy trì hoạt động trong suốt ván thi.",
        29: "F07 Phát câu hỏi đồng bộ: Game Worker tạo session_id, đặt deadline bằng System.nanoTime() rồi phát QUESTION. Các Player nhận cùng câu; chỉ công bố correct_answer khi câu đã đóng.",
        30: "F08 Nhận và khóa đáp án: ANSWER có ID yêu cầu, ván, câu và A/B/C/D. Server kiểm tra Player, trạng thái, deadline và nộp lặp; trả ANSWER_ACK chấp nhận hoặc từ chối.",
        31: "F09 Chấm điểm: Hết hạn, Server chấm một lần: đúng +100, còn lại +0. Không thưởng tốc độ do độ trễ Wi-Fi; Client chỉ hiển thị điểm từ Server.",
        38: "Hình 1 Kiến trúc Java Client-Server trong mạng LAN",
        39: "Nguyên tắc: Server giữ trạng thái chuẩn, câu hỏi, thời hạn, đáp án đúng và điểm. Admin App là Client có quyền quản trị, có thể chạy cùng máy Server hoặc máy khác trong LAN; không truy cập SQLite trực tiếp.",
        43: "Quy tắc chuyển trạng thái: Game Worker là nơi duy nhất cập nhật ván. Mặc định 10 câu, 15 giây/câu; countdown và kết quả đều 3 giây. Start cần ít nhất 2 Player đang kết nối, tất cả ready và đủ câu hợp lệ. Chờ hết hạn dù mọi người đã trả lời; reset xóa điểm và ready.",
        45: "Kênh truyền: TCP Socket, cổng LAN mặc định 5050, cấu hình được. JSON UTF-8 được serialize một dòng và kết thúc bằng byte newline. Receiver ghép byte trong buffer, chỉ parse dòng hoàn chỉnh; giới hạn mỗi message 64 KiB. Không dùng Java object serialization.",
        57: "1.  Admin kết nối TCP tới IP LAN và cổng Server, xác thực, cấu hình phòng và CRUD câu hỏi qua lệnh có kiểm tra quyền.",
        64: "Accept thread nhận kết nối bằng ServerSocket. Mỗi connection có reader đọc InputStream và writer duy nhất ghi OutputStream qua queue riêng có giới hạn; Client chậm không chặn broadcast cả phòng.",
        65: "Một Game Worker đọc BlockingQueue sự kiện, quản lý trạng thái và JDBC connection; handler không tự thay đổi điểm. Worker xếp message vào queue gửi riêng, không ghi socket trực tiếp.",
        66: "Deadline dùng System.nanoTime() tại Server; đồng hồ Client chỉ hiển thị. Desktop cập nhật UI qua Platform.runLater trên JavaFX Application Thread; kết nối và đọc/ghi chặn chạy ở worker riêng.",
        72: "5.7 Triển khai và vận hành trong LAN",
        73: "Mô hình: Chọn một máy chạy Server console; các máy cùng switch hoặc Wi-Fi chạy Admin và Player. Ví dụ Server 192.168.1.10, Client nhập 192.168.1.10:5050. Khi thử trên một máy, dùng 127.0.0.1.",
        74: "Kết nối: Server bind 0.0.0.0 hoặc IP LAN cụ thể; Client dùng IP thật của máy Server, không nhập 0.0.0.0. Kiểm tra firewall inbound TCP, đúng subnet và Wi-Fi client isolation. Không cần public IP hoặc mở cổng router.",
        75: "Vận hành: Máy Server không được ngủ hoặc tắt khi thi. Dừng Server sẽ ngắt mọi Client; restart đưa phòng về WAITING, không khôi phục ván dang dở. Sao lưu SQLite khi Server đã dừng ghi; dữ liệu đã commit giữ qua restart.",
        76: "Giới hạn bảo vệ: Chỉ triển khai trên LAN tin cậy với tài khoản và dữ liệu demo. Mật khẩu Admin lưu dạng băm có salt; TCP không mã hóa gói tin. Không dùng mật khẩu cá nhân hoặc công bố cổng dịch vụ ra Internet.",
        78: "Admin App: Kết nối và đăng nhập; màn hình phòng có danh sách ready, cấu hình, Start/Reset/Abort; ngân hàng gồm bảng, lọc và form; lịch sử kết quả chung cuộc. Điều khiển ván độc lập với tiến trình Server console.",
        79: "Player App: Kết nối bằng IP LAN, cổng, nickname và mã phòng; lobby; màn hình quiz với bốn nút đáp án lớn; kết quả từng câu và bảng hạng cuối. UI chỉ báo ghi nhận sau ACK, khóa nút khi gửi hoặc hết hạn.",
        80: "Thiết kế UI: JavaFX với CSS chung, tiếng Việt, font Segoe UI, màu nhấn xanh và trạng thái đi kèm chữ. Chuyển trang trong một cửa sổ, wrap/scroll nội dung dài và hỗ trợ resize; sự kiện mạng chuyển về UI qua Platform.runLater [3].",
        81: "Mã nguồn: Maven gồm nqb-shared, nqb-server, nqb-client-core, nqb-admin-client và nqb-player-client. Server chứa mạng, auth, Game Worker và repository; client-core không phụ thuộc UI. JavaFX chỉ nằm ở hai module desktop; test đặt trong src/test/java.",
        90: "Phối hợp: Thiện cung cấp Server và queue; Trường Thi phụ trách client-core và vòng câu/đáp án; Mỹ cung cấp repository và scoring. Cả nhóm thống nhất protocol, tích hợp hằng tuần, kiểm thử nhiều máy LAN và cùng giải thích được luồng TCP.",
        95: "Nghiệm thu: Tối thiểu ba Player chơi trọn ván trên ít nhất hai máy vật lý cùng LAN; có ít nhất 20 câu mẫu. Thử tăng tới 10 Player và ghi số liệu thực tế. Chỉ ghi Pass sau khi đã chạy kiểm thử, không suy ra từ build thành công.",
        97: "1.  Chạy Server trên một máy LAN, ghi IP/cổng; kiểm tra kết nối từ máy khác và kết nối Admin.",
        102: "6.  Kết thúc, đối chiếu ranking với lịch sử, reset chơi ván mới; restart Server để kiểm tra dữ liệu đã lưu.",
        103: "7.  Dùng Wireshark trên traffic demo được phép: handshake TCP, Follow TCP Stream và JSON phân tách newline. Chỉ dùng tài khoản thử nghiệm; giải thích TCP packet không tương đương một message.",
        108: "Mã nguồn tách Server, hai Client JavaFX và shared/client-core; IP, cổng và đường dẫn dữ liệu cấu hình được. Server không phụ thuộc JavaFX; Maven build và resources dùng UTF-8.",
        111: "Không đưa mật khẩu Admin hoặc dữ liệu thật vào repository; mật khẩu lưu dạng băm có salt. Cung cấp tài khoản demo riêng, chỉ sử dụng LAN tin cậy và không công bố cổng TCP ra Internet.",
        114: "[1] Java ServerSocket - https://docs.oracle.com/en/java/javase/21/docs/api/java.base/java/net/ServerSocket.html",
        115: "[2] Java BlockingQueue - https://docs.oracle.com/en/java/javase/21/docs/api/java.base/java/util/concurrent/BlockingQueue.html",
        116: "[3] OpenJFX Getting Started - https://openjfx.io/openjfx-docs/",
        117: "[4] Jackson databind - https://github.com/FasterXML/jackson-databind",
        118: "[5] Xerial SQLite JDBC - https://github.com/xerial/sqlite-jdbc",
        119: "[6] Wireshark User Guide - https://www.wireshark.org/docs/wsug_html_chunked/",
    }
    for i, text in updates.items():
        assert isinstance(blocks[i], Paragraph), i
        replace_paragraph(blocks[i], text, label=": " in text and text[:1].isalpha())
    replace_cell(tables[0].cell(3, 1), "Network Quiz Battle - Thi trắc nghiệm nhiều người bằng Java trong mạng LAN")
    replace_cell(tables[0].cell(4, 1), "Phạm Phúc Thiện - 2380610421 - F01, F02, F04, F05, F12; TCP Server, phòng, lỗi và điều phối nhóm.")
    replace_cell(tables[1].cell(1, 2), "Kết nối IP LAN/cổng bằng Java Socket; Server dùng ServerSocket.")
    replace_cell(tables[1].cell(1, 3), "IP LAN, cổng / trạng thái")
    replace_cell(tables[1].cell(1, 4), "Sai IP, firewall, timeout")
    replace_table(tables[2], [
        ["Admin App", "JavaFX; xác thực, CRUD, điều khiển phòng và xem lịch sử qua TCP."],
        ["Player App", "JavaFX; connect, lobby, ready, quiz và kết quả; dùng client-core chung."],
        ["TCP Server", "Java console; ServerSocket accept, reader/writer riêng; không phụ thuộc JavaFX."],
        ["Protocol và quyền", "Jackson, frame JSON UTF-8; quyền Admin và player_id gắn với connection."],
        ["Game Worker", "BlockingQueue điều phối ván, deadline và đáp án; một chủ sở hữu trạng thái."],
        ["Scoring", "Đúng +100, không thưởng tốc độ, đồng điểm đồng hạng; ranking từ Server."],
        ["Repository SQLite", "SQLite JDBC, PreparedStatement và transaction; chỉ Server truy cập DB."],
        ["Logging và config", "IP bind/cổng, data path, lỗi, giới hạn Client; vận hành trên máy LAN."],
    ])
    replace_table(tables[5], [
        ["Ngôn ngữ", "Java 21 LTS", "Dùng Java cho Server, Admin và Player; thống nhất JDK và UTF-8."],
        ["Giao diện", "JavaFX 21 và CSS", "Màn hình theo trạng thái Server; cập nhật trên JavaFX Application Thread."],
        ["Kết nối", "Socket, ServerSocket", "Tự lập trình accept/connect và InputStream/OutputStream; TCP trong LAN."],
        ["Định dạng", "Jackson, JSON UTF-8", "Buffer byte, frame newline; tối đa 64 KiB; không serialize object Java."],
        ["Đồng thời", "Thread, BlockingQueue", "Reader/writer mỗi connection; một Game Worker cập nhật trạng thái."],
        ["Thời gian", "System.nanoTime()", "Đo deadline tại Server; Instant dùng cho thời điểm lưu lịch sử."],
        ["Cơ sở dữ liệu", "SQLite JDBC", "File DB tại Server; PreparedStatement và transaction lưu kết quả."],
        ["Xác thực", "PBKDF2, SecureRandom", "Băm có salt và kiểm tra quyền; dùng tài khoản demo trên LAN tin cậy."],
        ["Triển khai", "Máy Server trong LAN", "IP LAN/cổng cấu hình; kiểm tra firewall và Wi-Fi; không cần Internet."],
        ["Kiểm thử", "JUnit 5, Wireshark", "Unit/integration Socket; quan sát TCP và test nhiều máy LAN."],
        ["Công cụ", "Maven, IDE Java, Git", "Chia 5 module, khóa dependency; build và resources UTF-8."],
    ])
    replace_cell(tables[6].cell(1, 1), "TCP/IP, Socket, ServerSocket")
    replace_cell(tables[6].cell(2, 2), "Demo một JSON chia qua nhiều read hoặc nhiều JSON trong một read.")
    replace_cell(tables[6].cell(6, 1), "SQLite JDBC và transaction")
    replace_cell(tables[6].cell(8, 1), "Địa chỉ IP và mạng LAN")
    replace_cell(tables[6].cell(8, 2), "Phân biệt loopback/IP LAN/bind; kiểm tra cổng và capture TCP.")
    replace_table(tables[7], [
        ["Tuần 1", "Thống nhất phạm vi Java, LAN và đề cương", "Hồ sơ gửi giảng viên", "Cả nhóm"],
        ["Tuần 2", "Chốt protocol, Maven, kiến trúc/UI; thử Socket trên hai máy LAN", "Schema, sơ đồ, thử cổng", "Thiện: mạng; Mỹ: DB; Trường Thi: UI"],
        ["Tuần 3", "Server/Client console, Admin auth, phòng và ba Client", "F01, F02, F04, F05", "Thiện; Trường Thi hỗ trợ Client"],
        ["Tuần 4", "SQLite JDBC, CRUD, Admin JavaFX và dữ liệu mẫu", "F03; nền tảng F09, F11", "Mỹ; cả nhóm tích hợp"],
        ["Tuần 5", "Game Worker, Player JavaFX, ready, câu hỏi, ACK và bảng điểm", "F06 - F11 end-to-end", "Trường Thi; Mỹ: điểm/lịch sử"],
        ["Tuần 6", "Heartbeat, lỗi mạng/queue, firewall và test nhiều máy LAN", "F12; biên bản kiểm thử", "Thiện chủ trì; cả nhóm"],
        ["Tuần 7", "Hoàn thiện UI, nhập liệu, tài liệu và số liệu test", "Bản demo và báo cáo", "Cả nhóm"],
        ["Tuần 8", "Tổng duyệt, quay video và bàn giao", "Mã nguồn, dữ liệu, README", "Cả nhóm"],
    ])
    replace_cell(tables[8].cell(1, 4), "Java TCP Server, quyền Admin, phòng, lỗi LAN và tích hợp.")
    replace_cell(tables[8].cell(2, 4), "Admin JavaFX, SQLite JDBC, CRUD, chấm điểm và lịch sử.")
    replace_cell(tables[8].cell(3, 4), "Player JavaFX, client-core, vòng câu/đáp án và bảng hạng.")
    replace_cell(tables[9].cell(1, 0), "Firewall, sai IP hoặc Wi-Fi cô lập Client")
    replace_cell(tables[9].cell(1, 2), "Thử hai máy sớm; kiểm tra IP LAN, inbound TCP và client isolation; chuẩn bị switch/LAN dự phòng.")
    replace_cell(tables[9].cell(2, 2), "Chốt nqb-shared và schema version 1; test framing và tích hợp mỗi tuần.")
    replace_cell(tables[10].cell(15, 1), "Ba Player trên ít nhất hai máy vật lý LAN")
    replace_cell(tables[10].cell(15, 2), "Chơi trọn ván qua IP LAN; score/ranking/history khớp")
    replace_cell(tables[10].cell(16, 1), "Sai IP/firewall; tăng tới 10 Player")
    replace_cell(tables[10].cell(16, 2), "Lỗi kết nối rõ, UI không treo; ghi số liệu tải thử thực tế")

    doc.core_properties.title = "Đề cương Network Quiz Battle bằng Java trong mạng LAN"
    doc.core_properties.author = "Nhóm VUA"
    doc.core_properties.keywords = "Java, JavaFX, TCP, Socket, ServerSocket, LAN, Maven, SQLite JDBC"
    buf = BytesIO()
    doc.save(buf)
    changed = {"word/document.xml", "docProps/core.xml", "word/media/image1.png"}
    with ZipFile(BytesIO(source_bytes)) as original, ZipFile(BytesIO(buf.getvalue())) as edited, ZipFile(TARGET, "w", ZIP_DEFLATED) as output:
        for part in original.infolist():
            if part.filename == "word/media/image1.png":
                data = diagram_bytes()
            elif part.filename in changed:
                data = edited.read(part.filename)
            else:
                data = original.read(part.filename)
            output.writestr(part, data)

    final = Document(TARGET)
    full_text = "\n".join(p.text for p in final.paragraphs) + "\n" + "\n".join(c.text for t in final.tables for r in t.rows for c in r.cells)
    for term in ["Python", "PySide6", "Qt", "TLS", "Serverix", "hosting", "online", "time.monotonic", ".py"]:
        assert term not in full_text, term
    for row in [1, 2, 5, 6]:
        assert final.tables[0].cell(row, 1).text == tables[0].cell(row, 1).text
    assert final.tables[12]._tbl.xml == tables[12]._tbl.xml
    assert final.tables[11]._tbl.xml == tables[11]._tbl.xml
    codes = [c.strip() for row in final.tables[8].rows[1:] for c in row.cells[3].text.split(",")]
    assert sorted(codes) == [f"F{i:02d}" for i in range(1, 13)]
    with ZipFile(BytesIO(source_bytes)) as a, ZipFile(TARGET) as b:
        assert a.namelist() == b.namelist()
        for part in a.namelist():
            if part not in changed:
                assert a.read(part) == b.read(part), part
    assert sha256(SOURCE.read_bytes()).hexdigest() == EXPECTED_HASH
    report = {"target": str(TARGET), "source_unchanged": True, "preserved_parts": len(manifest) - len(changed), "changed_parts": sorted(changed), "one_owner_per_function": True}
    (TMP / "structural_qa.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
