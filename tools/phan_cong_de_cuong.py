from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


ROOT = Path(r"E:\_HUTECH\DO_AN_LAP_TRINH_MANG")
SOURCE = ROOT / "output" / "documents" / "De_cuong_Network_Quiz_Battle.docx"
TARGET = ROOT / "output" / "documents" / "De_cuong_Network_Quiz_Battle_Da_phan_cong.docx"


def update_cell(cell, text, size=9.0, center=False):
    cell.text = text
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.08
        for run in paragraph.runs:
            run.font.name = "Arial"
            run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
            run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial")
            run.font.size = Pt(size)
            run.font.color.rgb = RGBColor(0, 0, 0)


def remove_empty_spacer_after(table):
    node = table._element.getnext()
    if node is not None and node.tag == qn("w:p") and not node.xpath(".//w:t"):
        node.getparent().remove(node)


doc = Document(SOURCE)

# Cập nhật phần thông tin nhóm bằng đúng tên và MSSV sinh viên đã điền.
group = doc.tables[0]
update_cell(
    group.cell(4, 1),
    "Phạm Phúc Thiện - 2380610421 - F01, F02, F04, F05, F12; Server và điều phối nhóm.",
    size=9.3,
)
update_cell(
    group.cell(5, 1),
    "Trịnh Viết Mỹ - 2380601382 - F03, F09, F11; câu hỏi, điểm và dữ liệu.\n"
    "Nguyễn Trường Thi - 2380602092 - F06, F07, F08, F10; Client và thi đấu.",
    size=9.3,
)

# Gắn kế hoạch với người chịu trách nhiệm thay cho các nhãn phân công chung.
plan = doc.tables[7]
plan_changes = {
    2: {
        3: "Thiện: giao thức; Mỹ: dữ liệu; Trường Thi: giao diện",
    },
    3: {
        1: "Server kết nối, đăng nhập, quản lý phòng; lớp dữ liệu câu hỏi",
        2: "F01, F02, F04, F05; nền tảng F03",
        3: "Thiện; Mỹ",
    },
    4: {
        1: "Ngân hàng câu hỏi, Client, sảnh chờ và luồng thi đấu",
        2: "F03, F06; giao diện F05, F07, F08",
        3: "Mỹ; Trường Thi",
    },
    5: {
        1: "Phát câu hỏi, nhận đáp án, điểm, bảng hạng, lịch sử và lỗi",
        2: "F07 - F12",
        3: "Thiện: F12; Mỹ: F09, F11; Trường Thi: F07, F08, F10",
    },
    6: {
        3: "Cả nhóm",
    },
    7: {
        3: "Cả nhóm",
    },
    8: {
        3: "Cả nhóm",
    },
}
for row_idx, columns in plan_changes.items():
    for col_idx, value in columns.items():
        update_cell(plan.cell(row_idx, col_idx), value, size=8.4, center=(col_idx == 0))

# Chia đủ F01-F12 cho ba người, mỗi chức năng có một chủ sở hữu chính.
assignment = doc.tables[8]
assignments = [
    (
        "Nhóm trưởng",
        "Phạm Phúc Thiện",
        "2380610421",
        "F01, F02, F04, F05, F12",
        "Server TCP, tài khoản, phòng thi, xử lý lỗi và tích hợp.",
    ),
    (
        "Thành viên",
        "Trịnh Viết Mỹ",
        "2380601382",
        "F03, F09, F11",
        "Câu hỏi và nhập CSV; SQLite, điểm và lịch sử trận.",
    ),
    (
        "Thành viên",
        "Nguyễn Trường Thi",
        "2380602092",
        "F06, F07, F08, F10",
        "Client Swing, sảnh chờ, câu hỏi, đáp án và bảng hạng.",
    ),
]
for row_idx, values in enumerate(assignments, start=1):
    for col_idx, value in enumerate(values):
        update_cell(assignment.cell(row_idx, col_idx), value, size=8.5, center=(col_idx in (2, 3)))

remove_empty_spacer_after(plan)
remove_empty_spacer_after(assignment)

doc.core_properties.comments = "Đã cập nhật phân công cụ thể cho ba thành viên trong đề cương."
doc.save(TARGET)

check = Document(TARGET)
all_codes = set()
for row in check.tables[8].rows[1:]:
    all_codes.update(code.strip() for code in row.cells[3].text.split(","))
expected = {f"F{i:02d}" for i in range(1, 13)}
assert all_codes == expected, (sorted(all_codes), sorted(expected))
assert "Phạm Phúc Thiện" in check.tables[0].cell(4, 1).text
assert "Trịnh Viết Mỹ" in check.tables[0].cell(5, 1).text
assert "Nguyễn Trường Thi" in check.tables[0].cell(5, 1).text
assert "Phamphucthien2325@gmail.com" in check.tables[0].cell(6, 1).text
print(TARGET)
